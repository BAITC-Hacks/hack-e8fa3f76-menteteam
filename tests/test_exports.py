from datetime import date
from io import BytesIO

from docx import Document

from core.models import MeetingResult, Segment, Task


def sample():
    return MeetingResult(title="Қазақша кеңес", summary="Айжан: есеп", participants={"SPEAKER_00": "Айжан"},
                         tasks=[Task(title="Есеп; жіберу", owner="Айжан", evidence="Айжан, есеп жібер",
                                     due_date=date(2026, 9, 24), status="Выполнено", email="a@example.test")],
                         transcript=[Segment(id="s", speaker="SPEAKER_00", text="Айжан, есеп жібер")])


def test_docx_contains_date_status_and_evidence():
    from exporters import export_docx
    document = Document(BytesIO(export_docx(sample())))
    cells = [c.text for table in document.tables for row in table.rows for c in row.cells]
    assert "2026-09-24" in cells
    assert "Выполнено" in cells
    assert any("Айжан, есеп жібер" in p.text for p in document.paragraphs)


def test_pdf_handles_long_unicode_tasks():
    from exporters import export_pdf
    meeting = sample()
    meeting.tasks[0].title = "Қазақша тапсырма " * 400
    assert export_pdf(meeting).startswith(b"%PDF")


def test_anonymization_removes_names_addresses_and_source():
    from exporters import anonymize
    meeting = sample()
    meeting.source_path = "/private/Айжан.wav"
    result = anonymize(meeting)
    assert "Айжан" not in result.model_dump_json()
    assert "a@example.test" not in result.model_dump_json()
    assert result.source_path == ""
    assert meeting.tasks[0].owner == "Айжан"


def test_calendar_escapes_text_and_uses_stable_uid():
    from exporters import export_ics
    meeting = sample()
    result = export_ics(meeting).decode().replace("\r\n ", "")
    assert "DTSTART;VALUE=DATE:20260924" in result
    assert "Есеп\\; жіберу" in result
    assert f"UID:{meeting.id}-{meeting.tasks[0].id}@alem.local" in result
