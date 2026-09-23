from io import BytesIO
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from core.models import MeetingResult


def _content(result: MeetingResult):
    rows = [["Поручение", "Ответственный", "Срок"]]
    rows.extend([[t.title, t.owner, t.deadline] for t in result.tasks])
    return rows


def export_docx(result: MeetingResult) -> bytes:
    doc = Document()
    doc.add_heading(f"Протокол: {result.title}", 0)
    doc.add_paragraph(f"Язык распознавания: {result.language}")
    doc.add_heading("Саммари", level=1); doc.add_paragraph(result.summary or "Не сформировано")
    doc.add_heading("Решения", level=1)
    for item in result.decisions: doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Поручения", level=1)
    table = doc.add_table(rows=1, cols=3); table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, _content(result)[0]): cell.text = value
    for values in _content(result)[1:]:
        for cell, value in zip(table.add_row().cells, values): cell.text = value
    for task in result.tasks:
        doc.add_paragraph(f"Evidence [{task.timestamp:.1f} сек]: {task.evidence}")
    doc.add_heading("Транскрипт", level=1)
    for s in result.transcript: doc.add_paragraph(f"[{s.start:.1f}–{s.end:.1f}] {s.speaker}: {s.text}")
    if result.questions:
        doc.add_heading("Нужно уточнить", level=1)
        for q in result.questions: doc.add_paragraph(q, style="List Bullet")
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def export_pdf(result: MeetingResult) -> bytes:
    """Generate a PDF and embed an installed Unicode font when available."""
    font_path = "/usr/share/fonts/TTF/DejaVuSans.ttf"
    if __import__("os").path.exists(font_path):
        try: pdfmetrics.registerFont(TTFont("DejaVu", font_path)); font_name = "DejaVu"
        except Exception: font_name = "Helvetica"
    else: font_name = "Helvetica"
    styles = getSampleStyleSheet(); styles["Normal"].fontName = font_name; styles["Title"].fontName = font_name; styles["Heading2"].fontName = font_name
    story = [Paragraph(f"Протокол: {result.title}", styles["Title"]), Spacer(1, 10),
             Paragraph("Саммари", styles["Heading2"]), Paragraph(result.summary or "Не сформировано", styles["Normal"]),
             Paragraph("Решения", styles["Heading2"])]
    for d in result.decisions: story.append(Paragraph(f"• {d}", styles["Normal"]))
    story.append(Paragraph("Поручения", styles["Heading2"]))
    table = Table(_content(result), repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#dce8f5")), ("GRID", (0,0), (-1,-1), .5, colors.grey), ("VALIGN", (0,0), (-1,-1), "TOP")]))
    story.extend([table, Paragraph("Транскрипт", styles["Heading2"])])
    for s in result.transcript: story.append(Paragraph(f"[{s.start:.1f}–{s.end:.1f}] {s.speaker}: {s.text}", styles["Normal"]))
    if result.questions:
        story.append(Paragraph("Нужно уточнить", styles["Heading2"]))
        for q in result.questions: story.append(Paragraph(f"• {q}", styles["Normal"]))
    stream = BytesIO(); SimpleDocTemplate(stream, pagesize=A4).build(story); return stream.getvalue()
